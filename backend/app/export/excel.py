"""Professional Excel export of query result datasets using openpyxl."""
from __future__ import annotations

import io
import re
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14, color="1F3A5F")
META_FONT = Font(italic=True, size=9, color="666666")
TOTAL_FONT = Font(bold=True)
TOTAL_FILL = PatternFill("solid", fgColor="E8EEF5")
THIN = Side(style="thin", color="C9D3DE")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

MONEY_HINTS = ("total", "amount", "value", "rate", "outstanding", "debit", "credit", "price", "cost", "net", "balance", "paid", "grand", "difference", "revenue", "sales", "purchases")
QTY_HINTS = ("qty", "quantity", "count", "invoices", "rows", "items", "employees", "bins", "units")
PCT_HINTS = ("pct", "percent", "share", "change", "growth", "margin")
NO_TOTAL = ("rate", "pct", "percent", "share", "change", "growth", "level", "id", "year", "month", "date", "balance_qty", "qty_after")
DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
DATETIME_RE = re.compile(r"^\d{4}-\d{2}-\d{2}[ T]\d{2}:\d{2}")


def _label(col: str) -> str:
    return re.sub(r"[_]+", " ", str(col)).strip().title()


def _number_format(col: str, sample: Any) -> str | None:
    name = str(col).lower()
    if any(h in name for h in PCT_HINTS):
        return '0.0"%"'
    if any(h in name for h in QTY_HINTS):
        return "#,##0.##"
    if any(h in name for h in MONEY_HINTS):
        return "#,##0.00"
    if isinstance(sample, float):
        return "#,##0.00"
    if isinstance(sample, int):
        return "#,##0"
    return None


def _coerce(value: Any) -> Any:
    if isinstance(value, str):
        if DATE_RE.match(value):
            try:
                return datetime.strptime(value, "%Y-%m-%d").date()
            except ValueError:
                return value
        if DATETIME_RE.match(value):
            try:
                return datetime.fromisoformat(value.replace(" ", "T")[:19])
            except ValueError:
                return value
    if isinstance(value, (dict, list)):
        return str(value)
    return value


def _sheet_title(title: str, used: set[str]) -> str:
    base = re.sub(r"[\[\]\*\?/\\:]", " ", title).strip()[:28] or "Data"
    name, i = base, 2
    while name in used:
        name = f"{base[:25]} {i}"
        i += 1
    used.add(name)
    return name


def _write_dataset(ws: Worksheet, dataset: dict[str, Any], currency: str) -> None:
    columns: list[str] = dataset.get("columns") or []
    rows: list[dict[str, Any]] = dataset.get("rows") or []
    if not columns and rows:
        columns = list(rows[0].keys())

    ws["A1"] = dataset.get("title") or "ERPNext Data"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Generated {datetime.now().strftime('%Y-%m-%d %H:%M')} by ERPNext AI Assistant" + (f" · Currency: {currency}" if currency else "")
    ws["A2"].font = META_FONT

    header_row = 4
    for c_idx, col in enumerate(columns, start=1):
        cell = ws.cell(row=header_row, column=c_idx, value=_label(col))
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[header_row].height = 22

    numeric_cols: dict[int, bool] = {}
    for r_idx, row in enumerate(rows, start=header_row + 1):
        for c_idx, col in enumerate(columns, start=1):
            value = _coerce(row.get(col))
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            cell.border = BORDER
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                fmt = _number_format(col, value)
                if fmt:
                    cell.number_format = fmt
                cell.alignment = Alignment(horizontal="right")
                numeric_cols[c_idx] = True
            elif isinstance(value, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm"
            elif isinstance(value, date):
                cell.number_format = "yyyy-mm-dd"
                cell.alignment = Alignment(horizontal="center")

    last_data_row = header_row + len(rows)
    if rows:
        total_row = last_data_row + 1
        ws.cell(row=total_row, column=1, value="Total").font = TOTAL_FONT
        for c_idx, col in enumerate(columns, start=1):
            cell = ws.cell(row=total_row, column=c_idx)
            cell.fill = TOTAL_FILL
            cell.border = BORDER
            cell.font = TOTAL_FONT
            name = str(col).lower()
            if numeric_cols.get(c_idx) and not any(h in name for h in NO_TOTAL):
                letter = get_column_letter(c_idx)
                cell.value = f"=SUBTOTAL(9,{letter}{header_row + 1}:{letter}{last_data_row})"
                cell.number_format = _number_format(col, 1.0) or "#,##0.00"
                cell.alignment = Alignment(horizontal="right")
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(columns))}{last_data_row}"
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    for c_idx, col in enumerate(columns, start=1):
        values = [str(_label(col))] + [str(r.get(col) if r.get(col) is not None else "") for r in rows[:500]]
        width = min(max(len(v) for v in values) + 3, 50)
        ws.column_dimensions[get_column_letter(c_idx)].width = max(width, 10)


def _write_summary(ws: Worksheet, datasets: list[dict[str, Any]], currency: str) -> None:
    ws["A1"] = "Summary"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Generated {datetime.now().strftime('%Y-%m-%d %H:%M')}" + (f" · Currency: {currency}" if currency else "")
    ws["A2"].font = META_FONT
    r = 4
    for ds in datasets:
        summary = ds.get("summary") or {}
        if not summary:
            continue
        ws.cell(row=r, column=1, value=ds.get("title") or "Dataset").font = Font(bold=True, color="1F3A5F")
        r += 1
        for k, v in summary.items():
            if isinstance(v, (dict, list)):
                v = ", ".join(f"{a}: {b}" for a, b in v.items()) if isinstance(v, dict) else ", ".join(map(str, v))
            ws.cell(row=r, column=1, value=_label(k)).border = BORDER
            cell = ws.cell(row=r, column=2, value=_coerce(v))
            cell.border = BORDER
            if isinstance(v, float):
                cell.number_format = "#,##0.00"
            elif isinstance(v, int) and not isinstance(v, bool):
                cell.number_format = "#,##0"
            r += 1
        r += 1
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 40


def build_workbook(datasets: list[dict[str, Any]], currency: str = "") -> bytes:
    """Return .xlsx bytes with one sheet per dataset plus a Summary sheet."""
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    with_rows = [d for d in datasets if d.get("rows")]
    for ds in with_rows:
        ws = wb.create_sheet(_sheet_title(ds.get("title") or "Data", used))
        _write_dataset(ws, ds, currency)
    if any(d.get("summary") for d in datasets):
        ws = wb.create_sheet(_sheet_title("Summary", used), 0 if not with_rows else None)
        _write_summary(ws, datasets, currency)
    if not wb.sheetnames:
        ws = wb.create_sheet("Data")
        ws["A1"] = "No data"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
