"""Parse uploaded files (xlsx/xls/csv/pdf/docx/txt) into text the model can read
and a compact statistical profile for spreadsheets."""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import pandas as pd

from app.core.errors import ValidationError

SUPPORTED = {".xlsx": "excel", ".xlsm": "excel", ".xls": "excel", ".csv": "csv", ".tsv": "csv", ".pdf": "pdf", ".docx": "docx", ".txt": "text", ".md": "text"}
MAX_CONTEXT_CHARS = 14000
SAMPLE_ROWS = 15


@dataclass
class ParsedFile:
    kind: str
    preview: str  # short one-line description for the UI
    context: str  # text given to the model
    sheets: dict[str, pd.DataFrame] = field(default_factory=dict)


def detect_kind(filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED:
        raise ValidationError(f"Unsupported file type '{ext}'. Allowed: {', '.join(sorted(SUPPORTED))}")
    return SUPPORTED[ext]


# ----------------------------------------------------------------- spreadsheets

def _clean_frame(df: pd.DataFrame) -> pd.DataFrame:
    df = df.dropna(how="all").dropna(axis=1, how="all")
    df.columns = [str(c).strip() if not str(c).startswith("Unnamed") else f"col_{i + 1}" for i, c in enumerate(df.columns)]
    # Promote first row to header if the header row looks unnamed and first row is text
    if all(c.startswith("col_") for c in df.columns) and len(df) > 1:
        first = df.iloc[0]
        if first.map(lambda v: isinstance(v, str)).mean() > 0.6:
            df = df.iloc[1:].reset_index(drop=True)
            df.columns = [str(v).strip() or f"col_{i + 1}" for i, v in enumerate(first)]
    for col in df.columns:
        if df[col].dtype == object:
            converted = pd.to_numeric(df[col].astype(str).str.replace(",", "", regex=False).str.strip(), errors="coerce")
            if converted.notna().sum() >= max(1, int(df[col].notna().sum() * 0.9)):
                df[col] = converted
    return df.reset_index(drop=True)


def load_spreadsheets(path: Path, kind: str) -> dict[str, pd.DataFrame]:
    sheets: dict[str, pd.DataFrame] = {}
    if kind == "csv":
        sep = "\t" if path.suffix.lower() == ".tsv" else None
        df = pd.read_csv(path, sep=sep, engine="python", encoding_errors="replace")
        sheets[path.stem] = _clean_frame(df)
    else:
        engine = "xlrd" if path.suffix.lower() == ".xls" else "openpyxl"
        book = pd.read_excel(path, sheet_name=None, engine=engine)
        for name, df in book.items():
            if df.empty:
                continue
            sheets[str(name)] = _clean_frame(df)
    if not sheets:
        raise ValidationError("The spreadsheet has no data.")
    return sheets


def profile_frame(name: str, df: pd.DataFrame, sample_rows: int = SAMPLE_ROWS) -> str:
    lines = [f"### Sheet: {name}", f"Rows: {len(df)}, Columns: {len(df.columns)}"]
    lines.append("Columns: " + ", ".join(f"{c} ({df[c].dtype})" for c in df.columns))
    numeric = df.select_dtypes(include="number")
    if not numeric.empty:
        sums = ", ".join(f"{c}={float(numeric[c].sum()):,.2f}" for c in numeric.columns[:15])
        lines.append("Column totals: " + sums)
    categorical = [c for c in df.columns if df[c].dtype == object and 1 < df[c].nunique() <= 25]
    for c in categorical[:3]:
        if not numeric.empty:
            top = df.groupby(c)[numeric.columns[0]].sum().sort_values(ascending=False).head(8)
            lines.append(f"Sum of {numeric.columns[0]} by {c}: " + ", ".join(f"{k}={float(v):,.2f}" for k, v in top.items()))
    try:
        head = df.head(sample_rows).to_markdown(index=False)
    except Exception:
        head = df.head(sample_rows).to_string(index=False)
    lines.append(f"First {min(sample_rows, len(df))} rows:")
    lines.append(head)
    return "\n".join(lines)


# ------------------------------------------------------------------- documents

def read_pdf(path: Path, max_pages: int = 25) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    parts = []
    for i, page in enumerate(reader.pages[:max_pages]):
        text = page.extract_text() or ""
        if text.strip():
            parts.append(f"--- Page {i + 1} ---\n{text.strip()}")
    if len(reader.pages) > max_pages:
        parts.append(f"[{len(reader.pages) - max_pages} more pages not shown]")
    return "\n".join(parts) or "[No extractable text; the PDF may be scanned images.]"


def read_docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for t_index, table in enumerate(document.tables):
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        parts.append(f"[Table {t_index + 1}]\n" + "\n".join(rows))
    return "\n".join(parts)


def read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


# ---------------------------------------------------------------------- entry

def parse_file(path: Path, filename: str) -> ParsedFile:
    kind = detect_kind(filename)
    if kind in ("excel", "csv"):
        sheets = load_spreadsheets(path, kind)
        sections = [profile_frame(name, df) for name, df in sheets.items()]
        context = f"## Uploaded file: {filename}\n" + "\n\n".join(sections)
        total_rows = sum(len(df) for df in sheets.values())
        preview = f"{len(sheets)} sheet(s), {total_rows} rows"
        parsed = ParsedFile(kind=kind, preview=preview, context=_truncate(context), sheets=sheets)
    else:
        if kind == "pdf":
            text = read_pdf(path)
        elif kind == "docx":
            text = read_docx(path)
        else:
            text = read_text(path)
        words = len(text.split())
        parsed = ParsedFile(kind=kind, preview=f"{words} words", context=_truncate(f"## Uploaded file: {filename}\n{text}"))
    return parsed


def _truncate(text: str, limit: int = MAX_CONTEXT_CHARS) -> str:
    if len(text) <= limit:
        return text
    return text[:limit] + f"\n[... truncated, {len(text) - limit} more characters. Use analyze_uploaded_file for exact figures.]"


def dataframe_to_records(df: pd.DataFrame, limit: int = 200) -> list[dict[str, Any]]:
    out = df.head(limit).copy()
    for col in out.columns:
        if pd.api.types.is_datetime64_any_dtype(out[col]):
            out[col] = out[col].dt.strftime("%Y-%m-%d")
    records = out.to_dict(orient="records")
    clean: list[dict[str, Any]] = []
    for r in records:
        clean.append({k: (None if (isinstance(v, float) and pd.isna(v)) else (round(v, 2) if isinstance(v, float) else v)) for k, v in r.items()})
    return clean


def save_bytes(data: bytes, dest: Path) -> None:
    dest.parent.mkdir(parents=True, exist_ok=True)
    with io.open(dest, "wb") as fh:
        fh.write(data)
